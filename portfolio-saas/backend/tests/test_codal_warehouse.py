import hashlib
import io
from unittest.mock import Mock, patch

import pytest
from openpyxl import Workbook
from rest_framework.test import APIClient

from marketdata.codal_classification import classify_announcement
from marketdata.codal_parsers import (
    ParsedDocument,
    category_reconciles,
    extract_typed_facts,
    parse_excel,
    parse_number,
)
from marketdata.codal_pipeline import (
    CodalBlockedStorage,
    artifact_path,
    configuration_error,
    download_artifact,
    store_artifact,
)
from marketdata.models import (
    CodalAnnouncement,
    CodalFact,
    CodalParsedTable,
    CodalReport,
    WorkflowRun,
)
from marketdata.workflows import WorkflowOutcome, redact

pytestmark = pytest.mark.django_db


@pytest.mark.parametrize(
    ("code", "title", "category"),
    [
        ("let11", "Material disclosure", 1),
        ("let6", "Financial statements", 2),
        ("let58", "Monthly production", 3),
        ("", "گزارش هیئت مدیره", 4),
        ("let56", "Auditor opinion", 5),
        ("let20", "Assembly decision", 6),
        ("let55", "Capital increase", 7),
        ("let8", "Investment portfolio", 8),
        ("let19", "Governance", 9),
        ("", "صورت مالی تلفیقی شرکت فرعی", 10),
        ("", "امیدنامه عرضه عمومی", 11),
    ],
)
def test_all_eleven_categories_are_classified(code, title, category):
    announcement = CodalAnnouncement(symbol="x", title=title, code=code)
    assert classify_announcement(announcement)["category"] == category


def test_excel_is_lossless_and_persian_numbers_become_typed_facts():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly"
    sheet.append(["محصول", "مبلغ فروش"])
    sheet.append(["A", "۱۲۳٬۴۵۶"])
    data = io.BytesIO()
    workbook.save(data)

    parsed = parse_excel(data.getvalue())
    extract_typed_facts(parsed, 3, "1405-05-01")

    assert parsed.tables[0]["rows"] == [["A", "۱۲۳٬۴۵۶"]]
    assert parsed.facts[0]["numeric_value"] == 123456
    assert parse_number("(۱٬۲۰۰٫۵)") == pytest.approx(-1200.5)


def test_ocr_publication_requires_category_reconciliation_and_threshold(settings):
    settings.CODAL_OCR_CONFIDENCE_THRESHOLD = 0.90
    parsed = ParsedDocument(text="مبلغ فروش ۱۰۰", confidence=0.8999, used_ocr=True)
    extract_typed_facts(parsed, 3)
    assert category_reconciles(parsed, 3)
    assert parsed.confidence < settings.CODAL_OCR_CONFIDENCE_THRESHOLD


def test_workflow_json_redacts_proxy_credentials_and_signed_urls():
    value = redact({
        "proxy": "https://user:secret@proxy.example:443",
        "download": "https://bucket.example/a?X-Amz-Signature=secret&part=1",
        "api_key": "secret",
    })
    assert "secret" not in str(value)
    assert value["api_key"] == "***"
    assert "user" not in value["proxy"]


def test_workflow_run_persists_one_terminal_summary():
    outcome = WorkflowOutcome("test_job", endpoint="fixture", symbol="X")
    outcome.finish(WorkflowRun.Outcome.SUCCESS, rows_received=2, rows_accepted=1)
    row = WorkflowRun.objects.get()
    assert row.rows_received == 2
    assert row.rows_accepted == 1


def _pdf_response():
    response = Mock()
    response.is_redirect = response.is_permanent_redirect = False
    response.status_code = 200
    response.headers = {"Content-Type": "application/pdf", "Content-Length": "12"}
    response.raise_for_status.return_value = None
    response.iter_content.return_value = [b"%PDF-fixture"]
    return response


def test_proxy_download_to_checksum_addressed_storage(settings, tmp_path):
    settings.CODAL_HTTP_PROXY = "https://proxy.example"
    settings.CODAL_MAX_ARTIFACT_BYTES = 1024
    settings.CODAL_STORAGE_DIR = str(tmp_path)

    with patch(
        "marketdata.codal_pipeline.requests.Session.get", return_value=_pdf_response()
    ) as get:
        _url, content_type, content = download_artifact("https://codal.ir/a.pdf", "pdf")
        key, checksum = store_artifact(content, content_type, "pdf")

    assert get.call_args.kwargs["proxies"]["https"] == settings.CODAL_HTTP_PROXY
    assert checksum in key
    stored = tmp_path / key
    assert stored.read_bytes() == b"%PDF-fixture"
    assert hashlib.sha256(stored.read_bytes()).hexdigest() == checksum


def test_storing_identical_bytes_twice_writes_one_file(settings, tmp_path):
    """Content addressing is the dedupe: the same document must not store twice."""
    settings.CODAL_STORAGE_DIR = str(tmp_path)

    first, checksum = store_artifact(b"%PDF-same", "application/pdf", "pdf")
    second, again = store_artifact(b"%PDF-same", "application/pdf", "pdf")

    assert (first, checksum) == (second, again)
    assert len(list(tmp_path.rglob("*.pdf"))) == 1


def test_storage_key_cannot_escape_the_storage_root(settings, tmp_path):
    """`s3_key` comes back from the database, so it is treated as untrusted."""
    settings.CODAL_STORAGE_DIR = str(tmp_path)

    with pytest.raises(CodalBlockedStorage):
        artifact_path("../../etc/passwd")


def test_only_the_proxy_is_required_once_storage_is_writable(settings, tmp_path):
    """Storage needed five S3 values; on disk it needs one writable directory."""
    settings.CODAL_STORAGE_DIR = str(tmp_path)
    settings.CODAL_HTTP_PROXY = ""
    assert configuration_error() == ["CODAL_HTTP_PROXY"]

    settings.CODAL_HTTP_PROXY = "https://proxy.example"
    assert configuration_error() == []


def test_report_apis_expose_metadata_and_default_to_latest_revision(make_user):
    original_announcement = CodalAnnouncement.objects.create(
        symbol="کاما", title="Statement", code="let6", date_publish="1404-01-01", time_publish="10:00:00"
    )
    original = CodalReport.objects.create(
        announcement=original_announcement, category=2, report_type="Statements", status="parsed", quality="validated"
    )
    correction_announcement = CodalAnnouncement.objects.create(
        symbol="کاما", title="Correction", code="let6-c", date_publish="1404-01-02", time_publish="10:00:00"
    )
    correction = CodalReport.objects.create(
        announcement=correction_announcement, category=2, report_type="Statements", status="parsed", quality="validated", revision_of=original
    )
    CodalFact.objects.create(report=original, fact_code="financial.revenue", numeric_value=1)
    CodalFact.objects.create(report=correction, fact_code="financial.revenue", numeric_value=2)

    client = APIClient()
    client.force_authenticate(user=make_user(tier="PRO"))
    announcements = client.get("/api/market/announcements/")
    facts = client.get("/api/market/facts/?symbol=کاما")
    detail = client.get(f"/api/market/reports/{correction.pk}/")

    assert announcements.status_code == facts.status_code == detail.status_code == 200
    assert announcements.json()[0]["report_id"] == correction.pk
    assert [row["numeric_value"] for row in facts.json()] == ["2.000000000000"]
    assert detail.json()["revision_of"] == original.pk
