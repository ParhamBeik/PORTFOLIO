"""Bounded Codal artifact extraction into MongoDB."""
import functools
import hashlib
from io import BytesIO

import requests
from django.conf import settings
from django.utils import timezone

from .models import CodalAnnouncement, CodalReport
from .quota import ARCHIVE, reserve_request

PARSER_VERSION = "text-v1"


@functools.lru_cache(maxsize=1)
def _client():
    """One pooled MongoClient for the worker process.

    MongoClient owns a connection pool and is designed to be created once; a new
    one per extraction opened a fresh pool for a single document and left the
    old one to be garbage-collected.
    """
    if not settings.MONGO_URI:
        raise RuntimeError("MongoDB is not configured.")
    from pymongo import MongoClient

    return MongoClient(settings.MONGO_URI, serverSelectionTimeoutMS=5000)


def _document_collection():
    return _client()[settings.MONGO_DATABASE]["codal_documents"]


def _artifact(announcement):
    for kind, url in (
        ("excel", announcement.link_excel),
        ("pdf", announcement.link_pdf),
        ("html", announcement.link),
        ("attachment", announcement.link_attachment),
    ):
        if url:
            return kind, url
    return "", ""


def _text(kind, body):
    if kind == "excel":
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(body), read_only=True, data_only=True)
        return "\n".join(
            "\t".join(str(value) for value in row if value is not None)
            for sheet in book.worksheets
            for row in sheet.iter_rows(values_only=True)
        )
    if kind == "pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(body)).pages)
    from bs4 import BeautifulSoup

    return BeautifulSoup(body, "lxml").get_text("\n", strip=True)


def extract(announcement_id):
    announcement = CodalAnnouncement.objects.get(pk=announcement_id)
    report, _ = CodalReport.objects.get_or_create(
        announcement=announcement, defaults={"category": announcement.category}
    )
    kind, url = _artifact(announcement)
    if not url:
        report.status = CodalReport.Status.UNSUPPORTED
        report.error_code = "artifact_url_missing"
        report.save(update_fields=["status", "error_code", "updated_at"])
        return {"status": report.status}
    report.status = CodalReport.Status.FETCHING
    report.save(update_fields=["status", "updated_at"])
    try:
        # Codal artifacts are megabyte-scale downloads from the same provider
        # the live price loop depends on. Going through the shared quota keeps a
        # backlog of documents from spending the budget live prices need; the
        # archive bucket is the right one because this is bulk backfill, not a
        # customer-facing read.
        reserve_request(ARCHIVE)
        response = requests.get(url, timeout=(10, 60))
        response.raise_for_status()
        text = _text(kind, response.content)
    except Exception as exc:
        report.status = CodalReport.Status.FAILED
        report.error_code = type(exc).__name__[:64]
        report.save(update_fields=["status", "error_code", "updated_at"])
        raise
    if not text.strip():
        report.status = CodalReport.Status.UNSUPPORTED
        report.error_code = "text_unavailable"
        report.save(update_fields=["status", "error_code", "updated_at"])
        return {"status": report.status}
    digest = hashlib.sha256(response.content).hexdigest()
    try:
        _document_collection().replace_one(
            {"announcement_id": announcement.pk, "parser_version": PARSER_VERSION},
            {
                "announcement_id": announcement.pk,
                "parser_version": PARSER_VERSION,
                "artifact": {"kind": kind, "url": url, "sha256": digest},
                "text": text[:2_000_000],
                "extracted_at": timezone.now(),
            },
            upsert=True,
        )
    except Exception as exc:
        report.status = CodalReport.Status.BLOCKED_STORAGE
        report.error_code = type(exc).__name__[:64]
        report.save(update_fields=["status", "error_code", "updated_at"])
        raise
    report.status = CodalReport.Status.PARSED
    report.quality = CodalReport.Quality.DEGRADED
    report.parser_version = PARSER_VERSION
    report.error_code = ""
    report.extracted_at = timezone.now()
    report.save()
    return {"status": report.status, "sha256": digest}
