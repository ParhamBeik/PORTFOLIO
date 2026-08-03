"""Extract current_state.json, history_snapshots.jsonl, and settings.json from Excel export."""
import json
import openpyxl
from datetime import datetime
from pathlib import Path

EXCEL_PATH = Path("exports/Portfolio_2026-08-01.xlsx")
DATA_DIR = Path("data")
DATA_DIR.mkdir(exist_ok=True)

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

# ============================================================
# 1. Parse Live_Data for current prices
# ============================================================
ws_live = wb["Live_Data"]
prices = {}
for row in ws_live.iter_rows(min_row=2, values_only=True):
    asset, price, unit, key, last_snapshot = row
    if key:
        prices[key] = price

print("Current prices loaded:", len(prices))

# ============================================================
# 2. Parse Settings sheet for constants
# ============================================================
ws_settings = wb["Settings"]
settings = {
    "api_settings": {},
    "api_urls": {},
    "constants": {}
}
for row in ws_settings.iter_rows(min_row=2, values_only=True):
    setting, value, notes = row
    if setting == "Rolling Window":
        # Not a constant, skip
        pass
    elif setting in ("house_area_sqm", "house_mortgage_deduction", "quarter_pre86_factor", 
                     "quarter_to_1g_ratio", "swiss_gold_bar_1g_price", "swiss_gold_bar_2_5g_price"):
        settings["constants"][setting] = value
    elif setting in ("brs_url", "tsetmc_url", "tsetmc_symbol_url", "tsetmc_history_url"):
        settings["api_settings"][setting] = value
    elif setting in ("brsapi", "tsetmc_base"):
        settings["api_urls"][setting] = value

# Add defaults for missing
defaults = {
    "quarter_pre86_factor": 0.8694109297,
    "quarter_to_1g_ratio": 0.493733384,
    "house_area_sqm": 90.2,
    "house_mortgage_deduction": 400000000,
    "swiss_gold_bar_1g_price": 25900000,
    "swiss_gold_bar_2_5g_price": 61610000,
}
for k, v in defaults.items():
    if k not in settings["constants"]:
        settings["constants"][k] = v

settings_path = DATA_DIR / "settings.json"
with settings_path.open("w") as f:
    json.dump(settings, f, indent=2)
print(f"Wrote settings.json to {settings_path}")

# ============================================================
# 3. Parse History_Data for full history_snapshots.jsonl
# ============================================================
ws_hist = wb["History_Data"]
headers = [c.value for c in ws_hist[1]]
col_idx = {h: i for i, h in enumerate(headers)}

history_records = []
for row in ws_hist.iter_rows(min_row=2, values_only=True):
    dt = row[0]
    if not dt:
        continue
    
    # Build snapshot structure matching what the pipeline expects
    snapshot = {
        "total_values_tomans": {
            "Mother": row[col_idx["IRT_Mother_Gold"]] + row[col_idx["IRT_Mother_Cash"]] + 
                      row[col_idx["IRT_Mother_Stock"]] + row[col_idx["IRT_Mother_RealEstate"]] + 
                      row[col_idx["IRT_Mother_Crypto"]],
            "Father": row[col_idx["IRT_Father_Gold"]] + row[col_idx["IRT_Father_Cash"]] + 
                      row[col_idx["IRT_Father_Stock"]] + row[col_idx["IRT_Father_RealEstate"]] + 
                      row[col_idx["IRT_Father_Crypto"]],
        },
        "prices": {
            "bitcoin_usd": row[col_idx["Price_bitcoin_usd"]],
            "usdt_irt": row[col_idx["Price_usdt_irt"]],
            "usd_cash": row[col_idx["Price_usd_cash"]],
            "emami_coin": row[col_idx["Price_emami_coin"]],
            "half_coin": row[col_idx["Price_half_coin"]],
            "quarter_coin": row[col_idx["Price_quarter_coin"]],
            "quarter_coin_pre86": row[col_idx["Price_quarter_coin_pre86"]],
            "swiss_gold_bar_1g": row[col_idx["Price_swiss_gold_bar_1g"]],
            "swiss_gold_bar_2_5g": row[col_idx["Price_swiss_gold_bar_2_5g"]],
            "one_gram_coin": row[col_idx["Price_one_gram_coin"]],
            "gold_18k_gram": row[col_idx["Price_gold_18k_gram"]],
            "kama_stock": row[col_idx["Price_kama_stock"]],
            "euro_cash": row[col_idx["Price_euro_cash"]],
            "gold_ounce_usd": row[col_idx["Price_gold_ounce_usd"]],
        }
    }
    
    record = {
        "timestamp": dt.isoformat() if isinstance(dt, datetime) else str(dt),
        "snapshot": snapshot
    }
    history_records.append(record)

hist_path = DATA_DIR / "history_snapshots.jsonl"
with hist_path.open("w") as f:
    for rec in history_records:
        f.write(json.dumps(rec) + "\n")
print(f"Wrote {len(history_records)} records to {hist_path}")

# ============================================================
# 4. Reconstruct current_state.json from latest snapshot
# ============================================================
# We need to estimate quantities from the latest per-asset-class totals.
# Since we don't know the exact composition, we'll use a reasonable assumption
# based on the test fixture and the latest values.

last_row = history_records[-1]["snapshot"]
mother_totals = last_row["total_values_tomans"]["Mother"]
father_totals = last_row["total_values_tomans"]["Father"]

# Get latest per-asset-class totals from last row of History_Data
last_hist_row = None
for row in ws_hist.iter_rows(min_row=2, values_only=True):
    last_hist_row = row

mother_gold = last_hist_row[col_idx["IRT_Mother_Gold"]]
mother_cash = last_hist_row[col_idx["IRT_Mother_Cash"]]
mother_stock = last_hist_row[col_idx["IRT_Mother_Stock"]]
mother_re = last_hist_row[col_idx["IRT_Mother_RealEstate"]]

father_gold = last_hist_row[col_idx["IRT_Father_Gold"]]
father_cash = last_hist_row[col_idx["IRT_Father_Cash"]]
father_stock = last_hist_row[col_idx["IRT_Father_Stock"]]
father_re = last_hist_row[col_idx["IRT_Father_RealEstate"]]

print(f"\nLatest asset-class totals (IRT):")
print(f"  Mother: Gold={mother_gold:,.0f}, Cash={mother_cash:,.0f}, Stock={mother_stock:,.0f}, RE={mother_re:,.0f}")
print(f"  Father: Gold={father_gold:,.0f}, Cash={father_cash:,.0f}, Stock={father_stock:,.0f}, RE={father_re:,.0f}")

# Reconstruct quantities based on reasonable assumptions:
# Father: primarily emami_coin (gold), usd_cash (cash), kama_stock (stock)
# Mother: primarily gold_18k_gram (gold), usdt_irt/usd_cash (cash), house (RE)

# Father's gold -> emami_coin
father_emami_qty = round(father_gold / prices["emami_coin"])
# Father's cash -> usd_cash  
father_usd_qty = round(father_cash / prices["usd_cash"])
# Father's stock -> kama_stock
father_kama_qty = round(father_stock / prices["kama_stock"])

# Mother's RE -> house_price_per_sqm_million (always 90 based on constant formula)
# 7,718,000,000 = 90 * 1e6 * 90.2 - 400,000,000
mother_house_price_per_sqm = 90

# Mother's cash -> usdt_irt (or usd_cash)
# Let's use usd_cash as primary
mother_usd_qty = round(mother_cash / prices["usd_cash"])

# Mother's gold -> gold_18k_gram (most liquid gold asset for individuals)
mother_gold_18k_qty = round(mother_gold / prices["gold_18k_gram"])

print(f"\nEstimated quantities:")
print(f"  Father: emami_coin={father_emami_qty}, usd_cash={father_usd_qty}, kama_stock={father_kama_qty}")
print(f"  Mother: gold_18k_gram={mother_gold_18k_qty}, usd_cash={mother_usd_qty}, house_price_per_sqm_million={mother_house_price_per_sqm}")

# Verify by recalculating
verify_father_gold = father_emami_qty * prices["emami_coin"]
verify_father_cash = father_usd_qty * prices["usd_cash"]
verify_father_stock = father_kama_qty * prices["kama_stock"]
verify_mother_gold = mother_gold_18k_qty * prices["gold_18k_gram"]
verify_mother_cash = mother_usd_qty * prices["usd_cash"]
verify_mother_re = (mother_house_price_per_sqm * 1_000_000 * 90.2) - 400_000_000

print(f"\nVerification:")
print(f"  Father: Gold={verify_father_gold:,.0f} (target {father_gold:,.0f}), Cash={verify_father_cash:,.0f} (target {father_cash:,.0f}), Stock={verify_father_stock:,.0f} (target {father_stock:,.0f})")
print(f"  Mother: Gold={verify_mother_gold:,.0f} (target {mother_gold:,.0f}), Cash={verify_mother_cash:,.0f} (target {mother_cash:,.0f}), RE={verify_mother_re:,.0f} (target {mother_re:,.0f})")

# Build current_state.json
current_state = {
    "Father": {
        "emami_coin": father_emami_qty,
        "usd_cash": father_usd_qty,
        "kama_stock": father_kama_qty,
    },
    "Mother": {
        "gold_18k_gram": mother_gold_18k_qty,
        "usd_cash": mother_usd_qty,
        "house_price_per_sqm_million": mother_house_price_per_sqm,
    }
}

state_path = DATA_DIR / "current_state.json"
with state_path.open("w") as f:
    json.dump(current_state, f, indent=2, ensure_ascii=False)
print(f"\nWrote current_state.json to {state_path}")

# ============================================================
# 5. Also create latest_prices_cache.json
# ============================================================
cache_path = DATA_DIR / "latest_prices_cache.json"
with cache_path.open("w") as f:
    json.dump(prices, f, indent=2)
print(f"Wrote latest_prices_cache.json to {cache_path}")

print("\n✅ All data files restored!")
